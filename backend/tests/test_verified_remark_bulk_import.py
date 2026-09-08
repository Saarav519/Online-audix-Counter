"""Bin verification in bulk, from a location + remark sheet.

Ticking hundreds of bins one dropdown at a time is not realistic on a real
audit. The same remarks can be uploaded as a two-column file — but the upload
must be exactly as strict as the dropdown about WHICH remarks are allowed,
while being forgiving about how they were typed.
"""
import asyncio
import io
import os
import uuid
from datetime import datetime, timezone

import pytest
import requests

from conftest import get_admin_password

MONGO_URL = os.environ.get("MONGO_URL", "")
DB_NAME = os.environ.get("DB_NAME", "")
needs_db = pytest.mark.skipif(not (MONGO_URL and DB_NAME), reason="MONGO_URL/DB_NAME not set")
pytestmark = needs_db

BINS = [f"BIN-{i:02d}" for i in range(1, 6)]


def _run(coro_fn):
    async def _wrapped():
        from motor.motor_asyncio import AsyncIOMotorClient
        cl = AsyncIOMotorClient(MONGO_URL)
        try:
            return await coro_fn(cl[DB_NAME])
        finally:
            cl.close()
    return asyncio.run(_wrapped())


def _hdr(u):
    return {"X-User-Id": u["id"], "X-Username": u.get("username", "")}


@pytest.fixture(scope="module")
def portal(base_url):
    return f"{base_url}/api/audit/portal"


@pytest.fixture(scope="module")
def admin(portal):
    r = requests.post(f"{portal}/login",
                      json={"username": "admin", "password": get_admin_password()}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["user"]


@pytest.fixture
def site(portal, admin):
    """Client with five bins in the master, all scanned so they appear on the sheet."""
    code = f"BK{uuid.uuid4().hex[:6].upper()}"
    cid = requests.post(f"{portal}/clients", json={
        "name": f"TEST bulk {code}", "code": code, "client_type": "warehouse"},
        headers=_hdr(admin), timeout=30).json()["client"]["id"]

    lm = "location_code,zone\n" + "".join(f"{b},Z1\n" for b in BINS)
    requests.post(f"{portal}/clients/{cid}/import-location-master",
                  files={"file": ("lm.csv", io.BytesIO(lm.encode()), "text/csv")},
                  headers=_hdr(admin), timeout=30)

    sid = requests.post(f"{portal}/sessions", json={
        "client_id": cid, "name": "Bulk Session", "variance_mode": "bin-wise",
        "start_date": "2026-01-01T00:00:00+00:00"},
        headers=_hdr(admin), timeout=30).json()["session"]["id"]
    csv = "location,barcode,qty\n" + "".join(f"{b},890000000{i},10\n" for i, b in enumerate(BINS))
    requests.post(f"{portal}/sessions/{sid}/import-expected",
                  files={"file": ("s.csv", io.BytesIO(csv.encode()), "text/csv")}, timeout=30)

    async def scan(db):
        await db.synced_locations.insert_many([{
            "session_id": sid, "location_name": b, "device_name": "D",
            "items": [{"barcode": f"890000000{i}", "quantity": 10, "product_name": "X"}],
            "total_items": 1, "total_quantity": 10, "is_empty": False,
            "synced_at": datetime.now(timezone.utc).isoformat()} for i, b in enumerate(BINS)])
    _run(scan)

    yield {"client_id": cid, "session_id": sid}
    requests.delete(f"{portal}/clients/{cid}", timeout=60)


def _upload(portal, admin, cid, data, name="v.csv"):
    return requests.post(f"{portal}/clients/{cid}/verified-remarks/import",
                         files={"file": (name, io.BytesIO(data), "text/csv")},
                         headers=_hdr(admin), timeout=30)


def _sheet(portal, sid):
    rows = requests.get(f"{portal}/reports/{sid}/bin-wise", timeout=30).json()["report"]
    return {r["location"]: r.get("verified_remark", "") for r in rows}


def test_a_csv_sets_the_remarks_on_the_sheet(portal, admin, site):
    r = _upload(portal, admin, site["client_id"],
                b"Location,Remark\nBIN-01,Verified - Correct\nBIN-02,Verified - Damaged\n")
    assert r.status_code == 200, r.text
    assert r.json()["applied"] == 2
    sheet = _sheet(portal, site["session_id"])
    assert sheet["BIN-01"] == "Verified – Correct"
    assert sheet["BIN-02"] == "Verified – Damaged"


def test_typing_a_plain_hyphen_still_matches(portal, admin, site):
    """The options carry an en-dash. Nobody types that in Excel, and a sheet
    that silently rejects every row would be useless."""
    r = _upload(portal, admin, site["client_id"],
                b"location,remark\nBIN-03,verified - recount done\n")
    assert r.json()["applied"] == 1
    assert _sheet(portal, site["session_id"])["BIN-03"] == "Verified – Recount Done"


def test_an_unknown_remark_is_refused_not_invented(portal, admin, site):
    """The dropdown is the only vocabulary; an upload must not widen it."""
    r = _upload(portal, admin, site["client_id"],
                b"location,remark\nBIN-01,Verified - Correct\nBIN-04,Kuch bhi likh diya\n")
    body = r.json()
    assert body["applied"] == 1
    assert body["rejected_remarks"] == 1
    assert "BIN-04" in body["rejected_examples"][0]
    assert _sheet(portal, site["session_id"])["BIN-04"] == ""


def test_a_blank_cell_leaves_the_bin_alone(portal, admin, site):
    """Blank means "I did not check this one", not "wipe what is there"."""
    _upload(portal, admin, site["client_id"], b"location,remark\nBIN-05,Verified - Correct\n")
    r = _upload(portal, admin, site["client_id"], b"location,remark\nBIN-05,\n")
    assert r.json()["applied"] == 0 and r.json()["cleared"] == 0
    assert _sheet(portal, site["session_id"])["BIN-05"] == "Verified – Correct"


def test_not_verified_clears_a_bin(portal, admin, site):
    _upload(portal, admin, site["client_id"], b"location,remark\nBIN-01,Verified - Correct\n")
    r = _upload(portal, admin, site["client_id"], b"location,remark\nBIN-01,Not Verified\n")
    assert r.json()["cleared"] == 1
    assert _sheet(portal, site["session_id"])["BIN-01"] == ""


def test_a_location_not_in_the_master_is_flagged(portal, admin, site):
    """Applied anyway — clients without a location master must still work — but
    counted, so a typo does not pass unnoticed."""
    r = _upload(portal, admin, site["client_id"],
                b"location,remark\nBIN-99,Verified - Correct\n")
    body = r.json()
    assert body["applied"] == 1
    assert body["unknown_locations"] == 1
    assert body["unknown_examples"] == ["BIN-99"]


def test_an_xlsx_works_too(portal, admin, site):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Location", "Verified Remark"])
    ws.append(["BIN-02", "Verified – Short Found"])
    buf = io.BytesIO()
    wb.save(buf)
    r = requests.post(
        f"{portal}/clients/{site['client_id']}/verified-remarks/import",
        files={"file": ("v.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_hdr(admin), timeout=30)
    assert r.status_code == 200, r.text
    assert r.json()["applied"] == 1
    assert _sheet(portal, site["session_id"])["BIN-02"] == "Verified – Short Found"


def test_the_import_is_logged_once_not_per_row(portal, admin, site):
    """Thousands of rows must not swamp the movement log; per-bin attribution
    still lives on the verified_remarks rows themselves."""
    cid = site["client_id"]
    _upload(portal, admin, cid,
            b"location,remark\nBIN-01,Verified - Correct\nBIN-02,Verified - Damaged\n"
            b"BIN-03,Verified - Correct\n")

    async def check(db):
        logs = await db.audit_logs.find(
            {"client_id": cid, "field_name": "verified_remark_bulk"}, {"_id": 0}).to_list(50)
        row = await db.verified_remarks.find_one({"client_id": cid, "location": "BIN-01"}, {"_id": 0})
        return logs, row
    logs, row = _run(check)
    assert len(logs) == 1, f"expected one summary entry, got {len(logs)}"
    assert "3 set" in logs[0]["new_value"]
    assert row["verified_by_username"] == admin["username"]
    assert row["verified_at"]


def test_the_template_lists_every_bin_and_its_current_remark(portal, admin, site):
    """An empty two-column file would just move the guesswork elsewhere."""
    from openpyxl import load_workbook
    cid = site["client_id"]
    _upload(portal, admin, cid, b"location,remark\nBIN-02,Verified - Damaged\n")

    r = requests.get(f"{portal}/clients/{cid}/verified-remarks/template", timeout=30)
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers.get("Content-Disposition", "")

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Verified Remarks"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Location", "Remark")
    filled = {loc: (rem or "") for loc, rem in rows[1:]}
    for b in BINS:
        assert b in filled, f"{b} missing from the template"
    assert filled["BIN-02"] == "Verified – Damaged", "an existing remark was not carried in"


def test_the_template_carries_an_excel_dropdown_of_the_valid_options(portal, admin, site):
    """The options carry an en-dash nobody types. A dropdown inside the sheet
    is what stops a hand-filled file failing on every row."""
    from openpyxl import load_workbook
    r = requests.get(f"{portal}/clients/{site['client_id']}/verified-remarks/template", timeout=30)
    wb = load_workbook(io.BytesIO(r.content))

    listed = [c[0] for c in wb["Valid Remarks"].iter_rows(min_row=2, values_only=True)]
    served = requests.get(f"{portal}/reports/{site['session_id']}/bin-wise",
                          timeout=30).json()["verified_remark_options"]
    assert listed == served, "template options drifted from what the report serves"

    dvs = wb["Verified Remarks"].data_validations.dataValidation
    assert dvs, "no dropdown on the Remark column"
    assert dvs[0].type == "list"
    assert "Valid Remarks" in dvs[0].formula1


def test_a_filled_in_template_uploads_back_cleanly(portal, admin, site):
    """The round trip is the whole point: download, fill, upload, done."""
    from openpyxl import load_workbook
    cid = site["client_id"]
    r = requests.get(f"{portal}/clients/{cid}/verified-remarks/template", timeout=30)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Verified Remarks"]

    wanted = {"BIN-01": "Verified – Correct", "BIN-03": "Verified – Short Found"}
    for row in ws.iter_rows(min_row=2):
        loc = row[0].value
        if loc in wanted:
            row[1].value = wanted[loc]
    buf = io.BytesIO()
    wb.save(buf)

    up = requests.post(f"{portal}/clients/{cid}/verified-remarks/import",
                       files={"file": ("filled.xlsx", buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                       headers=_hdr(admin), timeout=30)
    assert up.status_code == 200, up.text
    body = up.json()
    assert body["rejected_remarks"] == 0, body["rejected_examples"]
    assert body["unknown_locations"] == 0

    sheet = _sheet(portal, site["session_id"])
    assert sheet["BIN-01"] == "Verified – Correct"
    assert sheet["BIN-03"] == "Verified – Short Found"
